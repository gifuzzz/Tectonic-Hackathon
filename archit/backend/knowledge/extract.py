"""Turn Drive file content into plain text."""
from io import BytesIO

GOOGLE_EXPORTS = {
    "application/vnd.google-apps.document": "text/plain",
    "application/vnd.google-apps.spreadsheet": "text/csv",
    "application/vnd.google-apps.presentation": "text/plain",
}
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PLAIN_TYPES = {"application/json", "application/xml"}
MAX_DOWNLOAD_BYTES = 20 * 1024 * 1024


def is_downloadable_text(mime):
    return mime in (PDF, DOCX) or mime in PLAIN_TYPES or mime.startswith("text/")


def extract_text(service, item):
    """Return the text of a Drive file, or "" for unsupported or too-large files."""
    mime = item.get("mimeType", "")
    if mime in GOOGLE_EXPORTS:
        data = service.files().export(fileId=item["id"], mimeType=GOOGLE_EXPORTS[mime]).execute()
        return text_from_bytes(data, GOOGLE_EXPORTS[mime])
    if not is_downloadable_text(mime) or int(item.get("size", 0)) > MAX_DOWNLOAD_BYTES:
        return ""
    data = service.files().get_media(fileId=item["id"], supportsAllDrives=True).execute()
    return text_from_bytes(data, mime)


def text_from_bytes(data, mime):
    if mime == PDF:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if mime == DOCX:
        import docx

        document = docx.Document(BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8", errors="replace").lstrip("﻿")
